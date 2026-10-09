"""Uploaded CV file -> plain text (PDF, DOCX or TXT)."""
from __future__ import annotations

import io
from pathlib import Path


class CVTextError(ValueError):
    """The upload could not be turned into text. The message is safe to show the user."""


def extract_text(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf":
            import pdfplumber

            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif ext == ".docx":
            import docx

            document = docx.Document(io.BytesIO(content))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                parts.extend(cell.text for row in table.rows for cell in row.cells)
            text = "\n".join(parts)
        elif ext in {".txt", ".md"}:
            text = content.decode("utf-8", errors="replace")
        else:
            raise CVTextError("Upload a PDF, DOCX or TXT file.")
    except CVTextError:
        raise
    except ImportError:
        raise
    except Exception as exc:  # corrupt or password-protected files
        raise CVTextError(f"Could not read this {ext or 'file'}: {exc}") from exc

    if not text.strip():
        raise CVTextError("No text found in the file. Scanned images need OCR, which is not supported yet.")
    return text
